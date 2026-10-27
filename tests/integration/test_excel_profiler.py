"""Integration tests for Excel profiler."""

from collections.abc import Callable
from datetime import date, datetime, time
from pathlib import Path
from zipfile import ZipFile

import pytest
from openpyxl import Workbook, load_workbook

from ai_data_enrichment.domain.dataset_profile import PrimitiveType
from ai_data_enrichment.ingestion.excel_profiler import profile_worksheet
from ai_data_enrichment.ingestion.profiling_limits import ProfilingLimits

CellValue = str | int | float | bool | date | datetime | time | None
ExcelRows = list[list[CellValue]]


@pytest.fixture
def make_excel_workbook(
    tmp_path: Path,
) -> Callable[[ExcelRows], Path]:
    """Create a temporary Excel workbook with the supplied rows."""

    def _make(rows: ExcelRows) -> Path:
        workbook = Workbook()

        try:
            worksheet = workbook.active
            assert worksheet is not None
            worksheet.title = "Mountains"

            for row in rows:
                worksheet.append(row)

            path = tmp_path / "mountains.xlsx"
            workbook.save(path)

            return path
        finally:
            workbook.close()

    return _make


def test_excel_profiler_extracts_metadata_and_bounded_samples(
    make_excel_workbook: Callable[[ExcelRows], Path],
) -> None:
    """Profiler should extract worksheet metadata and limit sampled records."""

    path = make_excel_workbook(
        [
            ["Mountain", "Country", "height"],
            ["Everest", "Nepal", 8848.86],
            ["Mont Blanc", "France", 4808.73],
            ["Kilimanjaro", "Tanzania", 5895.00],
        ]
    )

    profile = profile_worksheet(
        path=path,
        sheet_name="Mountains",
        dataset_id="mountains",
        limits=ProfilingLimits(max_sample_records=2),
    )

    assert profile.dataset_id == "mountains"
    assert profile.name == "Mountains"
    assert profile.row_count == 3
    assert [column.name for column in profile.columns] == [
        "Mountain",
        "Country",
        "height",
    ]
    assert profile.sample_records == [
        {"Mountain": "Everest", "Country": "Nepal", "height": 8848.86},
        {"Mountain": "Mont Blanc", "Country": "France", "height": 4808.73},
    ]


def test_excel_profiler_counts_nulls_per_column(
    make_excel_workbook: Callable[[ExcelRows], Path],
) -> None:
    """Profiler should count missing values across all non-empty rows."""
    path = make_excel_workbook(
        [
            ["Mountain", "Country", "height"],
            ["Everest", "Nepal", 8848.86],
            ["Mont Blanc", None, 4808.73],
            [None, "Tanzania", 5895],
            ["Fuji", "Japan", None],
            [None, None, None],
        ]
    )

    profile = profile_worksheet(
        path=path,
        sheet_name="Mountains",
        dataset_id="mountains",
        limits=ProfilingLimits(max_sample_records=1),
    )

    assert profile.row_count == 4
    assert len(profile.sample_records) == 1
    assert [column.null_count for column in profile.columns] == [1, 1, 1]


def test_excel_profiler_infers_homogeneous_column_types(
    make_excel_workbook: Callable[[ExcelRows], Path],
) -> None:
    """Profiler should infer primitive types from column values."""
    path = make_excel_workbook(
        [
            ["Mountain", "Height", "Rank", "Active"],
            ["Everest", 8848.86, 1, True],
            ["Mont Blanc", 4808.73, 2, False],
            ["Kilimanjaro", 5895.25, 3, True],
        ]
    )

    profile = profile_worksheet(
        path=path,
        sheet_name="Mountains",
        dataset_id="mountains",
        limits=ProfilingLimits(max_sample_records=1),
    )

    assert [column.inferred_type for column in profile.columns] == [
        PrimitiveType.TEXT,
        PrimitiveType.NUMBER,
        PrimitiveType.INTEGER,
        PrimitiveType.BOOLEAN,
    ]


def test_excel_profiler_promotes_integer_and_float_to_number(
    make_excel_workbook: Callable[[ExcelRows], Path],
) -> None:
    """Profiler should infer Number for mixed integer and float values."""
    path = make_excel_workbook(
        [
            ["Mountain", "Height"],
            ["Everest", 8848],
            ["Mont Blanc", 4808.73],
            ["Kilimanjaro", 5895],
        ]
    )

    profile = profile_worksheet(
        path=path,
        sheet_name="Mountains",
        dataset_id="mountains",
        limits=ProfilingLimits(max_sample_records=1),
    )

    assert profile.columns[0].inferred_type == PrimitiveType.TEXT
    assert profile.columns[1].inferred_type == PrimitiveType.NUMBER


def test_excel_profiler_reports_mixed_column_types(
    make_excel_workbook: Callable[[ExcelRows], Path],
) -> None:
    """Profiler should report incompatible types with a warning."""
    path = make_excel_workbook(
        [
            ["Mountain", "Height"],
            ["Everest", 8848],
            ["Mont Blanc", 4808],
            ["Unknown", "not available"],
        ]
    )

    profile = profile_worksheet(
        path=path,
        sheet_name="Mountains",
        dataset_id="mountains",
        limits=ProfilingLimits(max_sample_records=1),
    )

    assert profile.columns[0].inferred_type == PrimitiveType.TEXT
    assert profile.columns[1].inferred_type == PrimitiveType.MIXED
    assert profile.columns[1].warnings == ["Mixed observed types: Integer, Text"]


@pytest.mark.parametrize(
    ("values", "expected_warning"),
    [
        (
            [8848, 4808.73, "unknown"],
            "Mixed observed types: Integer, Number, Text",
        ),
        (
            [True, 1, False],
            "Mixed observed types: Boolean, Integer",
        ),
    ],
)
def test_excel_profiler_reports_all_mixed_types(
    make_excel_workbook: Callable[[ExcelRows], Path],
    values: list[CellValue],
    expected_warning: str,
) -> None:
    """Profiler should report all incompatible observed types."""
    path = make_excel_workbook(
        [
            ["Mountain", "Height"],
            ["Everest", values[0]],
            ["Mont Blanc", values[1]],
            ["Unknown", values[2]],
        ]
    )

    profile = profile_worksheet(
        path=path,
        sheet_name="Mountains",
        dataset_id="mountains",
        limits=ProfilingLimits(max_sample_records=1),
    )

    assert profile.columns[0].inferred_type == PrimitiveType.TEXT
    assert profile.columns[1].inferred_type == PrimitiveType.MIXED
    assert profile.columns[1].warnings == [expected_warning]


def test_excel_profiler_infers_datetime_column(
    make_excel_workbook: Callable[[ExcelRows], Path],
) -> None:
    """Profiler should infer DateTime and serialize values as ISO 8601."""
    path = make_excel_workbook(
        [
            ["Mountain", "ObservedAt"],
            ["Everest", datetime(2024, 1, 2, 15, 45)],
            ["Mont Blanc", datetime(2024, 1, 3, 10, 30)],
        ]
    )

    profile = profile_worksheet(
        path=path,
        sheet_name="Mountains",
        dataset_id="mountains",
        limits=ProfilingLimits(max_sample_records=1),
    )

    assert profile.row_count == 2
    assert profile.columns[1].inferred_type == PrimitiveType.DATETIME

    assert profile.sample_records == [
        {
            "Mountain": "Everest",
            "ObservedAt": "2024-01-02T15:45:00",
        }
    ]


def test_excel_profiler_infers_date_column(
    make_excel_workbook: Callable[[ExcelRows], Path],
) -> None:
    """Profiler should infer Date and serialize values as ISO 8601."""
    path = make_excel_workbook(
        [
            ["Mountain", "FirstAscent"],
            ["Everest", date(1953, 5, 29)],
            ["Mont Blanc", date(1786, 8, 8)],
        ]
    )

    profile = profile_worksheet(
        path=path,
        sheet_name="Mountains",
        dataset_id="mountains",
        limits=ProfilingLimits(max_sample_records=1),
    )

    assert profile.row_count == 2
    assert profile.columns[1].inferred_type == PrimitiveType.DATE
    assert profile.sample_records == [
        {
            "Mountain": "Everest",
            "FirstAscent": "1953-05-29",
        }
    ]


def test_excel_profiler_preserves_midnight_datetime(
    make_excel_workbook: Callable[[ExcelRows], Path],
) -> None:
    """Profiler should not classify midnight DateTime values as Date."""
    path = make_excel_workbook(
        [
            ["Mountain", "ObservedAt"],
            ["Everest", datetime(2024, 1, 2, 0, 0)],
            ["Mont Blanc", datetime(2024, 1, 3, 0, 0)],
        ]
    )
    profile = profile_worksheet(
        path=path,
        sheet_name="Mountains",
        dataset_id="mountains",
        limits=ProfilingLimits(max_sample_records=1),
    )

    assert profile.columns[1].inferred_type == PrimitiveType.DATETIME

    assert profile.sample_records[0]["ObservedAt"] == ("2024-01-02T00:00:00")


def test_excel_profiler_collects_bounded_column_samples(
    make_excel_workbook: Callable[[ExcelRows], Path],
) -> None:
    """Profiler should collect bounded non-null samples per column."""
    path = make_excel_workbook(
        [
            ["Mountain", "Country"],
            ["Everest", "Nepal"],
            ["Mont Blanc", None],
            ["Kilimanjaro", "Tanzania"],
            ["Fuji", "Japan"],
        ]
    )
    profile = profile_worksheet(
        path=path,
        sheet_name="Mountains",
        dataset_id="mountains",
        limits=ProfilingLimits(max_sample_records=2),
    )

    assert profile.columns[0].sample_values == [
        "Everest",
        "Mont Blanc",
    ]

    assert profile.columns[1].sample_values == [
        "Nepal",
        "Tanzania",
    ]

    assert profile.columns[1].null_count == 1
    assert len(profile.sample_records) == 2


def test_excel_profiler_truncates_long_sample_values(
    make_excel_workbook: Callable[[ExcelRows], Path],
) -> None:
    """Profiler should limit text lengths in sampled values."""
    path = make_excel_workbook(
        [
            ["Mountain", "Height"],
            ["Kilimanjaro", 5895],
            ["Fuji", 3776],
            ["Mont Blanc", 4808.73],
        ]
    )
    profile = profile_worksheet(
        path=path,
        sheet_name="Mountains",
        dataset_id="mountains",
        limits=ProfilingLimits(
            max_sample_records=2,
            max_value_length=5,
        ),
    )

    assert profile.row_count == 3

    assert profile.columns[0].sample_values == [
        "Kilim",
        "Fuji",
    ]

    assert profile.sample_records == [
        {"Mountain": "Kilim", "Height": 5895},
        {"Mountain": "Fuji", "Height": 3776},
    ]


def test_excel_profiler_respects_max_profile_bytes(
    make_excel_workbook: Callable[[ExcelRows], Path],
) -> None:
    """Profiler should keep the serialized profile within its byte limit."""
    long_text = "A" * 180

    path = make_excel_workbook(
        [
            ["ID", "Description"],
            [1, long_text],
            [2, long_text],
            [3, long_text],
            [4, long_text],
            [5, long_text],
        ]
    )
    full_profile = profile_worksheet(
        path=path,
        sheet_name="Mountains",
        dataset_id="descriptions",
        limits=ProfilingLimits(
            max_sample_records=5,
            max_profile_bytes=12_000,
        ),
    )
    bounded_profile = profile_worksheet(
        path=path,
        sheet_name="Mountains",
        dataset_id="descriptions",
        limits=ProfilingLimits(
            max_sample_records=5,
            max_profile_bytes=900,
        ),
    )

    assert len(full_profile.model_dump_json().encode("utf-8")) > 900

    assert len(bounded_profile.model_dump_json().encode("utf-8")) <= 900

    assert bounded_profile.row_count == 5

    assert [column.name for column in bounded_profile.columns] == [
        "ID",
        "Description",
    ]

    assert full_profile.schema_warnings == []

    assert bounded_profile.schema_warnings == ["Sample data truncated to fit max_profile_bytes"]


def test_excel_profiler_rejects_insufficient_profile_budget(
    make_excel_workbook: Callable[[ExcelRows], Path],
) -> None:
    """Profiler should reject a budget smaller than mandatory metadata."""
    path = make_excel_workbook(
        [
            ["Mountain", "Height"],
            ["Everest", 8848],
        ]
    )

    with pytest.raises(
        ValueError,
        match="Dataset profile metadata exceeds max_profile_bytes=1",
    ):
        profile_worksheet(
            path=path,
            sheet_name="Mountains",
            dataset_id="mountains",
            limits=ProfilingLimits(max_profile_bytes=1),
        )


def test_excel_profiler_produces_deterministic_profile(
    make_excel_workbook: Callable[[ExcelRows], Path],
) -> None:
    """Profiler should produce identical bounded profiles for identical input."""

    long_text = "A" * 180

    path = make_excel_workbook(
        [
            ["ID", "Description"],
            [1, long_text],
            [2, long_text],
            [3, long_text],
            [4, long_text],
            [5, long_text],
        ]
    )

    limits = ProfilingLimits(
        max_sample_records=5,
        max_profile_bytes=900,
    )

    first_profile = profile_worksheet(
        path=path,
        sheet_name="Mountains",
        dataset_id="descriptions",
        limits=limits,
    )
    second_profile = profile_worksheet(
        path=path,
        sheet_name="Mountains",
        dataset_id="descriptions",
        limits=limits,
    )

    assert first_profile.model_dump_json() == second_profile.model_dump_json()

    assert len(first_profile.model_dump_json().encode("utf-8")) <= 900

    assert first_profile.row_count == 5


@pytest.mark.parametrize(
    ("headers", "expected_error"),
    [
        (
            ["Mountain", None],
            "Invalid column name at position 2",
        ),
        (
            ["Mountain", ""],
            "Invalid column name at position 2",
        ),
        (
            ["Mountain", "Mountain"],
            "Duplicate column name: Mountain",
        ),
    ],
)
def test_excel_profiler_rejects_invalid_headers(
    make_excel_workbook: Callable[[ExcelRows], Path],
    headers: list[CellValue],
    expected_error: str,
) -> None:
    """Profiler should reject missing and duplicate column names."""
    path = make_excel_workbook(
        [
            headers,
            ["Everest", 8848],
        ]
    )

    with pytest.raises(ValueError, match=expected_error):
        profile_worksheet(
            path=path,
            sheet_name="Mountains",
            dataset_id="mountains",
        )


def test_excel_profiler_rejects_missing_worksheet(
    make_excel_workbook: Callable[[ExcelRows], Path],
) -> None:
    """Profiler should report a missing worksheet."""
    path = make_excel_workbook(
        [
            ["Mountain", "Height"],
            ["Everest", 8848],
        ]
    )

    with pytest.raises(
        ValueError,
        match="Worksheet 'Missing' not found",
    ):
        profile_worksheet(
            path=path,
            sheet_name="Missing",
            dataset_id="mountains",
        )


def test_excel_profiler_rejects_empty_worksheet(
    make_excel_workbook: Callable[[ExcelRows], Path],
) -> None:
    """Profiler should reject a worksheet without a header row."""
    path = make_excel_workbook([])

    with pytest.raises(
        ValueError,
        match="Worksheet has no header row",
    ):
        profile_worksheet(
            path=path,
            sheet_name="Mountains",
            dataset_id="mountains",
        )


def test_excel_profiler_accepts_header_only_worksheet(
    make_excel_workbook: Callable[[ExcelRows], Path],
) -> None:
    """Profiler should accept a worksheet with headers and no data rows."""
    path = make_excel_workbook(
        [
            ["Mountain", "Country", "Height"],
        ]
    )
    profile = profile_worksheet(
        path=path,
        sheet_name="Mountains",
        dataset_id="mountains",
    )

    assert profile.row_count == 0
    assert profile.sample_records == []

    assert [column.name for column in profile.columns] == [
        "Mountain",
        "Country",
        "Height",
    ]

    assert all(column.inferred_type == PrimitiveType.UNKNOWN for column in profile.columns)

    assert all(column.null_count == 0 for column in profile.columns)
    assert all(column.sample_values == [] for column in profile.columns)


def test_excel_profiler_rejects_corrupted_workbook(
    tmp_path: Path,
) -> None:
    """Profiler should report a corrupted Excel workbook."""
    path = tmp_path / "corrupted.xlsx"
    path.write_bytes(b"This is not a valid XLSX file")

    with pytest.raises(
        ValueError,
        match="Invalid Excel workbook",
    ):
        profile_worksheet(
            path=path,
            sheet_name="Mountains",
            dataset_id="mountains",
        )


def test_excel_profiler_rejects_non_excel_zip(
    tmp_path: Path,
) -> None:
    """Profiler should reject a valid ZIP without Excel structure."""
    path = tmp_path / "fake_excel.xlsx"

    with ZipFile(path, "w") as archive:
        archive.writestr("readme.txt", "This is not an Excel workbook")

    with pytest.raises(
        ValueError,
        match="Invalid Excel workbook",
    ):
        profile_worksheet(
            path=path,
            sheet_name="Mountains",
            dataset_id="mountains",
        )


def test_excel_profiler_rejects_unsupported_file_format(
    tmp_path: Path,
) -> None:
    """Profiler should reject unsupported file formats."""
    path = tmp_path / "mountains.csv"
    path.write_text(
        "Mountain,Height\nEverest,8848\n",
        encoding="utf-8",
    )

    with pytest.raises(
        ValueError,
        match="Unsupported Excel file format",
    ):
        profile_worksheet(
            path=path,
            sheet_name="Mountains",
            dataset_id="mountains",
        )


def test_excel_profiler_reports_missing_workbook(
    tmp_path: Path,
) -> None:
    """Profiler should report a missing workbook without creating it."""
    path = tmp_path / "missing.xlsx"

    assert not path.exists()

    with pytest.raises(
        FileNotFoundError,
        match="missing.xlsx",
    ):
        profile_worksheet(
            path=path,
            sheet_name="Mountains",
            dataset_id="mountains",
        )

    assert not path.exists()


@pytest.mark.parametrize(
    ("filename", "target_column"),
    [
        ("mountains.xlsx", "height"),
        ("capitals.xlsx", "distance"),
    ],
)
def test_excel_profiler_profiles_sample_workbooks(
    filename: str,
    target_column: str,
) -> None:
    """Profiler should process the project's sample Excel workbooks."""
    path = Path(__file__).resolve().parents[2] / "samples" / filename

    assert path.is_file()

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet_name = workbook.sheetnames[0]
        worksheet = workbook[sheet_name]
        expected_headers = [cell.value for cell in next(worksheet.iter_rows(min_row=1, max_row=1))]
    finally:
        workbook.close()

    profile = profile_worksheet(
        path=path,
        sheet_name=sheet_name,
        dataset_id=path.stem,
    )

    assert profile.dataset_id == path.stem
    assert profile.name == sheet_name
    assert profile.row_count is not None
    assert profile.row_count > 0

    assert [column.name for column in profile.columns] == expected_headers

    assert len(profile.sample_records) <= 5
    assert len(profile.model_dump_json().encode("utf-8")) <= 12_000

    assert profile.row_count == 10

    target = next(column for column in profile.columns if column.name == target_column)

    assert target.inferred_type == PrimitiveType.UNKNOWN
    assert target.null_count == profile.row_count
    assert target.sample_values == []
    assert target.warnings == ["Column contains no non-null values; type cannot be inferred"]

    assert all(record[target_column] is None for record in profile.sample_records)


def test_excel_profiler_detects_excel_error_cells(
    make_excel_workbook: Callable[[ExcelRows], Path],
) -> None:
    """Profiler should identify Excel error cells rather than ordinary text."""
    path = make_excel_workbook(
        [
            ["Result"],
            ["#DIV/0!"],
        ]
    )

    profile = profile_worksheet(
        path=path,
        sheet_name="Mountains",
        dataset_id="excel-errors",
    )

    column = profile.columns[0]

    assert column.inferred_type == PrimitiveType.UNKNOWN
    assert any("Excel error" in warning for warning in column.warnings)


def test_excel_profiler_excludes_excel_errors_from_type_inference(
    make_excel_workbook: Callable[[ExcelRows], Path],
) -> None:
    """Excel errors should not affect inferred types or valid samples."""
    path = make_excel_workbook(
        [
            ["Result"],
            [42],
            ["#DIV/0!"],
            [43],
        ]
    )

    profile = profile_worksheet(
        path=path,
        sheet_name="Mountains",
        dataset_id="excel-errors-mixed",
    )

    column = profile.columns[0]

    assert profile.row_count == 3
    assert column.inferred_type == PrimitiveType.INTEGER
    assert column.null_count == 0
    assert column.sample_values == [42, 43]
    assert column.warnings == ["Excel error cells: 1"]


def test_excel_profiler_handles_time_only_values(
    make_excel_workbook: Callable[[ExcelRows], Path],
) -> None:
    """Profiler should preserve time-only values with a type warning."""
    path = make_excel_workbook(
        [
            ["StartTime"],
            [time(12, 30)],
            [time(14, 45)],
        ]
    )

    profile = profile_worksheet(
        path=path,
        sheet_name="Mountains",
        dataset_id="time-values",
    )

    column = profile.columns[0]

    assert profile.row_count == 2
    assert column.inferred_type == PrimitiveType.UNKNOWN
    assert column.null_count == 0
    assert column.sample_values == ["12:30:00", "14:45:00"]

    assert profile.sample_records == [
        {"StartTime": "12:30:00"},
        {"StartTime": "14:45:00"},
    ]

    assert any("time-only" in warning.lower() for warning in column.warnings)


def test_excel_profiler_detects_mixed_time_and_datetime(
    make_excel_workbook: Callable[[ExcelRows], Path],
) -> None:
    """Profiler should detect mixed time-only and datetime values."""
    path = make_excel_workbook(
        [
            ["EventTime"],
            [time(12, 30)],
            [datetime(2026, 10, 27, 14, 45)],
        ]
    )

    profile = profile_worksheet(
        path=path,
        sheet_name="Mountains",
        dataset_id="mixed-time-values",
    )

    column = profile.columns[0]

    assert column.inferred_type == PrimitiveType.MIXED
    assert column.null_count == 0

    assert column.sample_values == [
        "12:30:00",
        "2026-10-27T14:45:00",
    ]

    assert any("Mixed observed types" in warning for warning in column.warnings)

    assert any("time-only" in warning.lower() for warning in column.warnings)


def test_excel_profiler_warns_about_all_null_column(
    make_excel_workbook: Callable[[ExcelRows], Path],
) -> None:
    """Profiler should warn when a column contains only null values."""
    path = make_excel_workbook(
        [
            ["Mountain", "Height"],
            ["Everest", None],
            ["K2", None],
        ]
    )

    profile = profile_worksheet(
        path=path,
        sheet_name="Mountains",
        dataset_id="all-null-column",
    )

    column = profile.columns[1]

    assert profile.row_count == 2
    assert column.inferred_type == PrimitiveType.UNKNOWN
    assert column.null_count == 2
    assert column.sample_values == []

    assert column.warnings == ["Column contains no non-null values; type cannot be inferred"]


def test_excel_profiler_rejects_malformed_worksheet_xml(
    make_excel_workbook: Callable[[ExcelRows], Path],
    tmp_path: Path,
) -> None:
    """Profiler should report malformed worksheet XML."""
    valid_path = make_excel_workbook(
        [
            ["Mountain"],
            ["Everest"],
        ]
    )

    corrupted_path = tmp_path / "malformed.xlsx"

    with ZipFile(valid_path) as source, ZipFile(corrupted_path, "w") as target:
        for member in source.infolist():
            content = source.read(member.filename)

            if member.filename == "xl/worksheets/sheet1.xml":
                assert b"</sheetData>" in content
                content = content.replace(
                    b"</sheetData>",
                    b"</invalid>",
                )

            target.writestr(member, content)

    with pytest.raises(
        ValueError,
        match="Invalid Excel workbook",
    ):
        profile_worksheet(
            path=corrupted_path,
            sheet_name="Mountains",
            dataset_id="malformed",
        )


def test_excel_profiler_preserves_typed_temporal_values(
    make_excel_workbook: Callable[[ExcelRows], Path],
) -> None:
    """Text limits should not truncate typed ISO temporal values."""
    path = make_excel_workbook(
        [
            ["Date", "Time", "DateTime", "Description"],
            [
                date(2026, 10, 27),
                time(12, 30),
                datetime(2026, 10, 27, 14, 45),
                "Everest",
            ],
        ]
    )

    profile = profile_worksheet(
        path=path,
        sheet_name="Mountains",
        dataset_id="temporal-values",
        limits=ProfilingLimits(max_value_length=3),
    )

    assert profile.sample_records == [
        {
            "Date": "2026-10-27",
            "Time": "12:30:00",
            "DateTime": "2026-10-27T14:45:00",
            "Description": "Eve",
        }
    ]
