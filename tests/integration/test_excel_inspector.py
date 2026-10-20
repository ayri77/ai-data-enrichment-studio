"""Integration tests for read-only Excel inspection."""

import hashlib
import json
from pathlib import Path

import pytest
from openpyxl import Workbook
from typer.testing import CliRunner

from ai_data_enrichment.cli import app
from ai_data_enrichment.ingestion.excel_inspector import inspect_workbook


def _sha256(path: Path) -> str:
    """Hash file bytes to detect unintended input mutations."""
    return hashlib.sha256(path.read_bytes()).hexdigest()


@pytest.mark.integration
def test_inspector_handles_two_sheets_and_preserves_file(tmp_path: Path) -> None:
    """Inspection reads headers and data without changing the original file."""
    path = tmp_path / "example.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.title = "Mountains"
    worksheet.append(["Mountain", "Country", "height"])
    worksheet.append(["Everest", "Nepal", None])
    worksheet.append(["Mont Blanc", "France", None])
    second = workbook.create_sheet("Meta")
    second.append(["Key", "Value"])
    second.append(["version", 1])
    workbook.save(path)

    before = _sha256(path)
    summary = inspect_workbook(path)

    assert _sha256(path) == before
    assert len(summary.worksheets) == 2
    assert summary.worksheets[0].columns == ("Mountain", "Country", "height")
    assert summary.worksheets[0].row_count == 2
    assert summary.worksheets[1].row_count == 1


@pytest.mark.integration
def test_inspect_cli_outputs_valid_json(tmp_path: Path) -> None:
    """The CLI exposes a stable JSON representation for workbook summaries."""
    path = tmp_path / "capitals.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.append(["Capital_From", "Country_From", "Capital_To", "Country_To", "distance"])
    sheet.append(["Berlin", "Germany", "Vienna", "Austria", None])
    workbook.save(path)

    result = CliRunner().invoke(app, ["inspect", str(path)])
    assert result.exit_code == 0
    parsed = json.loads(result.stdout)
    assert parsed["worksheets"][0]["row_count"] == 1
    assert parsed["worksheets"][0]["columns"][-1] == "distance"


def test_invalid_extension_is_rejected(tmp_path: Path) -> None:
    """Unsupported files do not reach the workbook parser."""
    path = tmp_path / "example.csv"
    path.write_text("a,b\n1,2\n", encoding="utf-8")
    with pytest.raises(ValueError, match="Only .xlsx"):
        inspect_workbook(path)
