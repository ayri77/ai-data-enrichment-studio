"""Read-only XLSX inspection for the initial project bootstrap.

This is intentionally limited to worksheet names, first-row headers, and
non-empty data row counts. Full DatasetProfile modeling follows in M1.
"""

from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook


@dataclass(frozen=True, slots=True)
class WorksheetSummary:
    """Small, serialization-friendly summary of a worksheet."""

    name: str
    columns: tuple[str, ...]
    row_count: int


@dataclass(frozen=True, slots=True)
class WorkbookSummary:
    """Read-only summary of an XLSX workbook."""

    worksheets: tuple[WorksheetSummary, ...]


def inspect_workbook(path: Path) -> WorkbookSummary:
    """Inspect an XLSX file using its first row as the column header."""
    if path.suffix.lower() != ".xlsx":
        raise ValueError("Only .xlsx files are supported")
    if not path.is_file():
        raise FileNotFoundError(f"Workbook not found: {path}")

    summaries: list[WorksheetSummary] = []
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        for worksheet in workbook.worksheets:
            rows = worksheet.iter_rows(values_only=True)
            header = next(rows, ())
            columns = tuple(str(value).strip() if value is not None else "" for value in header)
            row_count = sum(
                1 for row in rows if any(value is not None and str(value).strip() for value in row)
            )
            summaries.append(
                WorksheetSummary(name=worksheet.title, columns=columns, row_count=row_count)
            )
    finally:
        workbook.close()
    return WorkbookSummary(worksheets=tuple(summaries))
