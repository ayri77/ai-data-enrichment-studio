"""Profile an Excel worksheet."""

from datetime import date, datetime, time
from pathlib import Path
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.cell.read_only import EmptyCell
from openpyxl.styles.numbers import is_datetime
from openpyxl.utils.exceptions import InvalidFileException
from pydantic import JsonValue

from ai_data_enrichment.domain.dataset_profile import (
    ColumnProfile,
    DatasetProfile,
    PrimitiveType,
)
from ai_data_enrichment.ingestion.profiling_limits import ProfilingLimits


def _to_json_value(
    value: object,
    number_format: str | None = None,
    *,
    max_value_length: int,
) -> JsonValue:
    """Convert an Excel cell value to a bounded JSON-compatible value."""
    if isinstance(value, datetime):
        if is_datetime(number_format) == "date":
            return value.date().isoformat()
        return value.isoformat()

    if isinstance(value, date):
        return value.isoformat()

    if isinstance(value, time):
        return value.isoformat()

    if isinstance(value, str):
        return value[:max_value_length]

    if value is None or isinstance(value, (bool, int, float)):
        return value

    raise ValueError(f"Unsupported Excel cell type: {type(value).__name__}")


def _infer_value_type(
    value: object,
    number_format: str | None = None,
) -> PrimitiveType:
    """Infer the canonical primitive type of an Excel cell."""
    if value is None:
        return PrimitiveType.UNKNOWN

    # bool is an int subclass, so infer it before numeric types.
    if isinstance(value, bool):
        return PrimitiveType.BOOLEAN

    if isinstance(value, int):
        return PrimitiveType.INTEGER

    if isinstance(value, float):
        return PrimitiveType.NUMBER

    if isinstance(value, str):
        return PrimitiveType.TEXT

    if isinstance(value, datetime):
        if is_datetime(number_format) == "date":
            return PrimitiveType.DATE
        return PrimitiveType.DATETIME

    if isinstance(value, date):
        return PrimitiveType.DATE

    return PrimitiveType.UNKNOWN


def _resolve_column_type(
    observed_types: set[PrimitiveType],
) -> PrimitiveType:
    """Resolve observed value types into a column type."""
    if not observed_types:
        return PrimitiveType.UNKNOWN

    if len(observed_types) == 1:
        return next(iter(observed_types))

    # Integer and floating-point observations share the Number type.
    if observed_types == {PrimitiveType.INTEGER, PrimitiveType.NUMBER}:
        return PrimitiveType.NUMBER

    return PrimitiveType.MIXED


def _enforce_profile_size(
    profile: DatasetProfile,
    max_bytes: int,
) -> DatasetProfile:
    """Reduce sampled data until the serialized profile fits the byte limit."""

    def serialized_size() -> int:
        return len(profile.model_dump_json().encode("utf-8"))

    if serialized_size() <= max_bytes:
        return profile

    profile.schema_warnings.append("Sample data truncated to fit max_profile_bytes")

    while serialized_size() > max_bytes:
        trimmed = False

        # Remove column samples first.
        for column in profile.columns:
            if column.sample_values:
                column.sample_values.pop()
                trimmed = True

                if serialized_size() <= max_bytes:
                    return profile

        if trimmed:
            continue

        # Preserve record relationships as long as possible.
        if profile.sample_records:
            profile.sample_records.pop()
            continue

        # Mandatory metadata cannot be reduced safely.
        raise ValueError(f"Dataset profile metadata exceeds max_profile_bytes={max_bytes}")

    return profile


def profile_worksheet(
    path: Path,
    sheet_name: str,
    dataset_id: str,
    limits: ProfilingLimits | None = None,
) -> DatasetProfile:
    """Profile an Excel worksheet without modifying its contents."""

    effective_limits = limits if limits is not None else ProfilingLimits()

    try:
        # Read cached formula results without evaluating formulas or changing the file.
        workbook = load_workbook(path, read_only=True, data_only=True)

    except InvalidFileException as exc:
        raise ValueError(f"Unsupported Excel file format: {path}") from exc

    except (BadZipFile, KeyError) as exc:
        raise ValueError(f"Invalid Excel workbook: {path}") from exc

    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"Worksheet '{sheet_name}' not found")

        worksheet = workbook[sheet_name]
        rows = worksheet.iter_rows(values_only=False)

        # Read and validate column headers.
        raw_header = next(rows, ())

        if not raw_header:
            raise ValueError("Worksheet has no header row")

        headers: list[str] = []

        for index, cell in enumerate(raw_header, start=1):
            name = cell.value

            if not isinstance(name, str) or not name.strip():
                raise ValueError(f"Invalid column name at position {index}")

            if name in headers:
                raise ValueError(f"Duplicate column name: {name}")

            headers.append(name)

        # Initialize profiling state.
        row_count = 0
        null_counts = [0] * len(headers)
        time_only_counts = [0] * len(headers)
        excel_error_counts = [0] * len(headers)
        sample_records: list[dict[str, JsonValue]] = []
        observed_types: list[set[PrimitiveType]] = [set() for _ in headers]
        column_samples: list[list[JsonValue]] = [[] for _ in headers]

        # Process all data rows.
        for row in rows:
            # Entirely empty rows do not count toward records or null statistics.
            if all(cell.value is None for cell in row):
                continue

            if len(row) > len(headers):
                raise ValueError(f"Row has {len(row)} columns, but header has {len(headers)}")

            if len(row) < len(headers):
                row = (
                    *row,
                    *(EmptyCell() for _ in range(len(headers) - len(row))),
                )
            row_count += 1

            # Count missing values in each column.
            for index, cell in enumerate(row):
                value = cell.value

                if value is None:
                    null_counts[index] += 1

                # Keep Excel error cells out of ordinary type inference.
                elif cell.data_type == "e":
                    excel_error_counts[index] += 1

                else:
                    observed_types[index].add(_infer_value_type(value, cell.number_format))
                    if isinstance(value, time):
                        time_only_counts[index] += 1
                    if len(column_samples[index]) < effective_limits.max_sample_records:
                        column_samples[index].append(
                            _to_json_value(
                                value,
                                cell.number_format,
                                max_value_length=effective_limits.max_value_length,
                            )
                        )

            # Collect only the configured number of sample records.
            if len(sample_records) < effective_limits.max_sample_records:
                record = {
                    name: _to_json_value(
                        cell.value,
                        cell.number_format,
                        max_value_length=effective_limits.max_value_length,
                    )
                    for name, cell in zip(headers, row, strict=True)
                }
                sample_records.append(record)

        # Build column profiles after processing all rows.
        column_profiles: list[ColumnProfile] = []

        for index, name in enumerate(headers):
            types = observed_types[index]
            inferred_type = _resolve_column_type(types)

            warnings: list[str] = []

            if inferred_type == PrimitiveType.MIXED:
                type_names = sorted(t.value for t in types)
                warnings.append(f"Mixed observed types: {', '.join(type_names)}")

            if row_count > 0 and null_counts[index] == row_count:
                warnings.append("Column contains no non-null values; type cannot be inferred")

            if excel_error_counts[index] > 0:
                warnings.append(f"Excel error cells: {excel_error_counts[index]}")

            if time_only_counts[index] > 0:
                warnings.append(
                    f"Time-only values without a canonical type: {time_only_counts[index]}"
                )

            column_profiles.append(
                ColumnProfile(
                    name=name,
                    inferred_type=inferred_type,
                    null_count=null_counts[index],
                    sample_values=column_samples[index],
                    warnings=warnings,
                )
            )

        profile = DatasetProfile(
            dataset_id=dataset_id,
            name=sheet_name,
            row_count=row_count,
            columns=column_profiles,
            sample_records=sample_records,
        )

        return _enforce_profile_size(
            profile,
            effective_limits.max_profile_bytes,
        )

    except (BadZipFile, KeyError, ParseError) as exc:
        raise ValueError(f"Invalid Excel workbook: {path}") from exc

    finally:
        workbook.close()
